"""Import the clinic's existing patient records from a CSV or Excel (.xlsx) file.

Dry run by default: validates every row and prints a report, writes nothing.
    python import_patients.py patients.xlsx               # check only
    python import_patients.py patients.csv --commit       # actually import
    python import_patients.py old.csv --date-format %m/%d/%Y --errors bad_rows.csv

Columns are matched by name (case/spaces ignored). Recognised headers:
    first_name, last_name  — or a single "name" / "full_name" column (split on the last space)
    phone (required), date_of_birth / dob / birthday, gender / sex, email, address, medical_notes / notes / allergies
Rows that fail validation, or match an existing patient (same name + phone), are skipped and reported.
"""

import argparse
import csv
import re
import sys
from datetime import datetime
from pathlib import Path

from pydantic import ValidationError
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Patient
from app.schemas import PatientCreate

ALIASES = {
    "firstname": "first_name", "given_name": "first_name",
    "lastname": "last_name", "surname": "last_name", "family_name": "last_name",
    "name": "full_name", "patient_name": "full_name",
    "phone_number": "phone", "mobile": "phone", "tel": "phone", "telephone": "phone",
    "dob": "date_of_birth", "birthday": "date_of_birth", "birth_date": "date_of_birth",
    "sex": "gender",
    "e_mail": "email", "email_address": "email",
    "notes": "medical_notes", "allergies": "medical_notes", "medical_history": "medical_notes",
}


def _key(header) -> str:
    k = re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")
    return ALIASES.get(k, k)


def read_rows(path: Path) -> list[dict]:
    """Rows as {normalised_header: value}; values are str, or date/datetime from Excel."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        headers = [_key(h) for h in next(rows)]
        return [dict(zip(headers, r)) for r in rows]
    with path.open(newline="", encoding="utf-8-sig") as f:  # utf-8-sig: Excel's CSV export adds a BOM
        return [{_key(k): v for k, v in r.items()} for r in csv.DictReader(f)]


def to_patient_data(row: dict, date_format: str) -> dict:
    data = {k: (v.strip() if isinstance(v, str) else v) for k, v in row.items() if v not in (None, "")}
    if "full_name" in data and not ("first_name" in data or "last_name" in data):
        first, _, last = str(data.pop("full_name")).rpartition(" ")
        data["first_name"], data["last_name"] = (first, last) if first else (last, "-")
    dob = data.get("date_of_birth")
    if isinstance(dob, datetime):
        data["date_of_birth"] = dob.date()
    elif isinstance(dob, str) and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", dob):
        data["date_of_birth"] = datetime.strptime(dob, date_format).date()
    if isinstance(data.get("phone"), (int, float)):  # Excel number cell; leading 0 is lost
        data["phone"] = str(int(data["phone"]))
    return {k: v for k, v in data.items() if k in PatientCreate.model_fields}


def run(path: Path, commit: bool, date_format: str, errors_path: Path | None) -> dict:
    rows = read_rows(path)
    errors, valid = [], []
    with SessionLocal() as db:
        seen = {(p.first_name.lower(), p.last_name.lower(), p.phone) for p in db.scalars(select(Patient))}
        for line, row in enumerate(rows, start=2):  # line 1 = header
            if not any(v not in (None, "") for v in row.values()):
                continue
            try:
                patient = PatientCreate(**to_patient_data(row, date_format))
            except (ValidationError, ValueError) as e:
                msg = "; ".join(f"{'.'.join(map(str, x['loc']))}: {x['msg']}" for x in e.errors()) \
                    if isinstance(e, ValidationError) else str(e)
                errors.append((line, msg, row))
                continue
            key = (patient.first_name.lower(), patient.last_name.lower(), patient.phone)
            if key in seen:
                errors.append((line, "duplicate (same name and phone already exists)", row))
                continue
            seen.add(key)
            valid.append(patient)

        if commit and valid:
            db.add_all(Patient(**p.model_dump()) for p in valid)
            db.commit()

    for line, msg, _ in errors:
        print(f"  line {line}: {msg}")
    if errors_path and errors:
        with errors_path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["line", "error", *errors[0][2].keys()])
            w.writerows([line, msg, *row.values()] for line, msg, row in errors)
    action = "Imported" if commit else "Would import (dry run, use --commit)"
    print(f"{action}: {len(valid)}   Skipped: {len(errors)}   Rows read: {len(rows)}")
    return {"valid": len(valid), "skipped": len(errors)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--commit", action="store_true", help="write to the database (default: dry run)")
    ap.add_argument("--date-format", default="%d/%m/%Y", help="for non-ISO dates (default %%d/%%m/%%Y)")
    ap.add_argument("--errors", type=Path, help="write skipped rows to this CSV")
    args = ap.parse_args()
    if args.file.suffix.lower() not in (".csv", ".xlsx"):
        sys.exit("Only .csv and .xlsx are supported (save old .xls files as .xlsx).")
    run(args.file, args.commit, args.date_format, args.errors)
